from __future__ import annotations

import warnings
from pathlib import Path


def load_workbook_quiet(path: str | Path, **kwargs):
    from openpyxl import load_workbook

    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Workbook contains no default style, apply openpyxl's default",
            category=UserWarning,
            module=r"openpyxl\.styles\.stylesheet",
        )
        return load_workbook(path, **kwargs)
