from __future__ import annotations

from collections import defaultdict
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from irbis_control.core import matcher as _matcher
from irbis_control.core.matcher import (
    SOURCE_FOREIGN_AGENTS,
    SOURCE_SUBSTANCES,
    CancelCallback,
    ProgressCallback,
    _cancelled,
    _extract_subfield,
    extract_isbns,
    normalize_author,
    normalize_title,
)
from irbis_control.core.models import ComparisonSummary, DatabaseRecord, MatchResult
from irbis_control.core.text import safe_text


def _publication_text(record: DatabaseRecord | None) -> str:
    if not record:
        return ""
    parts: list[str] = []
    for item in record.publication:
        city = _extract_subfield(item, "A")
        publisher = _extract_subfield(item, "C")
        year = _extract_subfield(item, "D")
        publication = ", ".join(value for value in (city, publisher, year) if value)
        if publication:
            parts.append(publication)
    return " | ".join(dict.fromkeys(parts))


def _unique_join(values: Iterable[Any], separator: str = " | ") -> str:
    cleaned = [safe_text(value) for value in values if safe_text(value)]
    return separator.join(dict.fromkeys(cleaned))


def _matched_database_records(
    results: list[MatchResult],
    source_type: str | None = None,
) -> list[DatabaseRecord]:
    """Возвращает уникальные подтверждённые записи TXT для выбранного источника."""
    unique: dict[int, DatabaseRecord] = {}
    for result in results:
        if result.status != "Совпадение" or result.database is None:
            continue
        if source_type is not None and result.source_type != source_type:
            continue
        unique.setdefault(result.database.record_number, result.database)
    return [unique[key] for key in sorted(unique)]


COMMON_MATCH_HEADERS = [
    "№",
    "Автор",
    "Название",
    "ISBN",
    "Инвентарные номера",
    "Издание: город, издательство, год",
    "Номер записи в текстовой базе",
]

SUBSTANCE_MATCH_HEADERS = [
    *COMMON_MATCH_HEADERS,
    "Почему добавлено",
    "Файл-источник",
    "Лист",
    "Строка",
]

FOREIGN_AGENT_MATCH_HEADERS = [
    *COMMON_MATCH_HEADERS,
    "Совпавшее поле TXT",
    "Совпавшее значение",
    "№ в реестре",
    "Иностранный агент",
    "Тип иностранного агента",
    "Дата включения",
    "Вид записи в реестре",
    "Файл реестра",
    "Лист",
    "Строка",
]

COMBINED_MATCH_HEADERS = [
    *COMMON_MATCH_HEADERS,
    "Список совпадений",
    "Способ совпадения",
    "Совпавшее значение",
]

REVIEW_MATCH_HEADERS = [
    "Источник проверки",
    "Способ совпадения",
    "Точность, %",
    "Почему требуется проверка",
    "Совпавшее значение",
    "Автор записи",
    "Название записи",
    "Номер записи в базе",
    "Файл-источник",
    "Лист",
    "Строка",
    "Решение оператора",
]


def _confirmed_results_by_record(
    results: list[MatchResult],
    source_type: str | None = None,
) -> dict[int, list[MatchResult]]:
    grouped: dict[int, list[MatchResult]] = defaultdict(list)
    for result in results:
        if result.status != "Совпадение" or result.database is None:
            continue
        if source_type is not None and result.source_type != source_type:
            continue
        grouped[result.database.record_number].append(result)
    return grouped


def _report_records(
    results: list[MatchResult],
    source_type: str | None = None,
    *,
    deduplicate: bool = False,
    sort_by: str = "record",
) -> tuple[list[DatabaseRecord], dict[int, list[MatchResult]]]:
    """Готовит строки отчёта, не изменяя правила сопоставления и установки меток."""
    records = _matched_database_records(results, source_type)
    grouped = _confirmed_results_by_record(results, source_type)

    if deduplicate:
        unique: dict[tuple[Any, ...], DatabaseRecord] = {}
        merged: dict[int, list[MatchResult]] = defaultdict(list)
        for record in records:
            normalized_isbns = [isbn for value in record.isbns for isbn in extract_isbns(value)]
            if normalized_isbns:
                identity: tuple[Any, ...] = ("isbn", normalized_isbns[0])
            else:
                title = normalize_title(next(iter(record.titles), ""))
                author = normalize_author(next(iter(record.authors), ""))
                identity = ("title_author", title, author) if title or author else ("record", record.record_number)
            selected = unique.setdefault(identity, record)
            merged[selected.record_number].extend(grouped.get(record.record_number, []))
        records = list(unique.values())
        grouped = dict(merged)

    def sort_key(record: DatabaseRecord) -> tuple[Any, ...]:
        if sort_by == "title":
            value = normalize_title(next(iter(record.titles), ""))
        elif sort_by == "author":
            value = normalize_author(next(iter(record.authors), ""))
        elif sort_by == "isbn":
            value = next((isbn for raw in record.isbns for isbn in extract_isbns(raw)), "")
        else:
            return (record.record_number,)
        return (not bool(value), value, record.record_number)

    records.sort(key=sort_key)
    return records, grouped


def _parallel_join(values: Iterable[Any]) -> str:
    """Объединяет значения построчно, сохраняя соответствие между колонками."""
    return "\n".join(safe_text(value) for value in values)


def _base_txt_row(number: int, record: DatabaseRecord) -> list[Any]:
    return [
        number,
        _unique_join(record.authors),
        _unique_join(record.titles),
        _unique_join(record.isbns),
        _unique_join(record.inventory_numbers),
        _publication_text(record),
        record.record_number,
    ]


def _substance_match_row(
    number: int,
    record: DatabaseRecord,
    record_results: list[MatchResult],
) -> list[Any]:
    reasons = []
    for result in record_results:
        method = safe_text(result.method).strip()
        matched_value = safe_text(result.matched_value).strip()
        reasons.append(f"{method} — {matched_value}" if method and matched_value else method or matched_value)
    return [
        *_base_txt_row(number, record),
        _parallel_join(reasons),
        _parallel_join(Path(result.excel.source_file).name for result in record_results),
        _parallel_join(result.excel.sheet_name for result in record_results),
        _parallel_join(result.excel.row_number for result in record_results),
    ]


def _foreign_agent_field(result: MatchResult) -> str:
    if ":" in result.method:
        return result.method.rsplit(":", 1)[-1].strip()
    return result.method


def _foreign_agent_match_row(
    number: int,
    record: DatabaseRecord,
    record_results: list[MatchResult],
) -> list[Any]:
    entries = [result.foreign_agent for result in record_results]
    return [
        *_base_txt_row(number, record),
        _parallel_join(_foreign_agent_field(result) for result in record_results),
        _parallel_join(result.matched_value for result in record_results),
        _parallel_join(entry.registry_number if entry else "" for entry in entries),
        _parallel_join(entry.name if entry else "" for entry in entries),
        _parallel_join(entry.agent_type if entry else "" for entry in entries),
        _parallel_join(entry.inclusion_date if entry else "" for entry in entries),
        _parallel_join(result.note for result in record_results),
        _parallel_join(Path(entry.source_file).name if entry else "" for entry in entries),
        _parallel_join(entry.sheet_name if entry else "" for entry in entries),
        _parallel_join(entry.row_number if entry else "" for entry in entries),
    ]


def _combined_match_row(
    number: int,
    record: DatabaseRecord,
    record_results: list[MatchResult],
) -> list[Any]:
    return [
        *_base_txt_row(number, record),
        _parallel_join(result.source_type for result in record_results),
        _parallel_join(result.method for result in record_results),
        _parallel_join(result.matched_value for result in record_results),
    ]


def _style_match_sheet(
    worksheet,
    row_count: int,
    headers: list[str],
    widths: dict[str, int],
    row_fill_color: str,
) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    row_fill = PatternFill("solid", fgColor=row_fill_color)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(row_count, 1)}"
    worksheet.row_dimensions[1].height = 42

    for row_number, row in enumerate(worksheet.iter_rows(min_row=2, max_row=row_count), start=2):
        max_lines = max(
            (safe_text(cell.value).count("\n") + 1 for cell in row),
            default=1,
        )
        worksheet.row_dimensions[row_number].height = min(120, max(30, 18 * max_lines))
        for cell in row:
            cell.fill = row_fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width


def _add_results_sheet(
    workbook,
    *,
    title: str,
    headers: list[str],
    records: list[DatabaseRecord],
    grouped_results: dict[int, list[MatchResult]],
    row_builder,
    table_name: str,
    table_style: str,
    row_fill_color: str,
    widths: dict[str, int],
    active: bool = False,
    cancel_cb: CancelCallback | None = None,
):
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    worksheet = workbook.active if active else workbook.create_sheet()
    worksheet.title = title
    worksheet.append(headers)

    for number, record in enumerate(records, start=1):
        if number % 500 == 0:
            _cancelled(cancel_cb)
        worksheet.append(row_builder(number, record, grouped_results.get(record.record_number, [])))

    _style_match_sheet(
        worksheet,
        len(records) + 1,
        headers,
        widths,
        row_fill_color,
    )

    if records:
        last_column = get_column_letter(len(headers))
        table = Table(
            displayName=table_name,
            ref=f"A1:{last_column}{len(records) + 1}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name=table_style,
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)
    return worksheet


def _add_summary_sheet(workbook, summary: ComparisonSummary, *, active: bool = False):
    from openpyxl.styles import Font, PatternFill

    worksheet = workbook.active if active else workbook.create_sheet()
    worksheet.title = "Сводка"
    worksheet.append(["Показатель", "Значение"])
    rows = [
        ("Источник библиографических записей", summary.database_file),
        ("Проверено библиографических записей", summary.database_records),
        ("Строк в перечнях веществ", summary.excel_rows),
        ("Найдено строк перечней веществ", summary.matched_excel_rows),
        ("Записей с совпадениями по веществам", summary.substance_matched_records),
        ("Строк в реестре иностранных агентов", summary.foreign_agent_rows),
        ("Найдено строк реестра иностранных агентов", summary.matched_foreign_agent_rows),
        ("Записей с совпадениями по иностранным агентам", summary.foreign_agent_matched_records),
        ("Пограничных совпадений для ручной проверки", summary.review_rows),
    ]
    for row in rows:
        worksheet.append(row)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    worksheet.column_dimensions["A"].width = 48
    worksheet.column_dimensions["B"].width = 64
    worksheet.freeze_panes = "A2"
    return worksheet


def _add_review_sheet(
    workbook,
    results: list[MatchResult],
    *,
    active: bool = False,
    cancel_cb: CancelCallback | None = None,
):
    """Добавляет пограничные результаты, которые запрещено применять автоматически."""
    review_results = [
        result
        for result in results
        if result.status in {"Возможное совпадение", "Отклонено вручную"} and result.database is not None
    ]
    if not review_results:
        return None

    worksheet = workbook.active if active else workbook.create_sheet()
    worksheet.title = "Требует проверки"
    worksheet.append(REVIEW_MATCH_HEADERS)
    for number, result in enumerate(review_results, start=1):
        if number % 500 == 0:
            _cancelled(cancel_cb)
        record = result.database
        worksheet.append(
            [
                result.source_type,
                result.method,
                result.confidence,
                result.note or "Недостаточно данных для автоматической установки метки",
                result.matched_value,
                _unique_join(record.authors),
                _unique_join(record.titles),
                record.record_number,
                Path(result.excel.source_file).name,
                result.excel.sheet_name,
                result.excel.row_number,
                result.status,
            ]
        )

    _style_match_sheet(
        worksheet,
        len(review_results) + 1,
        REVIEW_MATCH_HEADERS,
        {
            "A": 24,
            "B": 38,
            "C": 14,
            "D": 62,
            "E": 48,
            "F": 34,
            "G": 52,
            "H": 24,
            "I": 34,
            "J": 22,
            "K": 12,
            "L": 24,
        },
        "FCE4D6",
    )
    return worksheet


# Принимает совпадения и сводку, формирует выбранные листы Excel и сохраняет отчёт атомарно.
def export_results(
    output_path: str | Path,
    results: list[MatchResult],
    summary: ComparisonSummary,
    progress_cb: ProgressCallback | None = None,
    cancel_cb: CancelCallback | None = None,
    report_options: dict[str, Any] | None = None,
) -> Path:
    """Создаёт выбранные пользователем листы с точными совпадениями."""
    from openpyxl import Workbook

    options = {
        "substances": True,
        "foreign_agents": True,
        "combined": False,
        "summary": False,
        "deduplicate": False,
        "sort": "record",
    }
    if report_options:
        for key in ("substances", "foreign_agents", "combined", "summary", "deduplicate"):
            if key in report_options:
                options[key] = bool(report_options[key])
        if report_options.get("sort") in {"record", "title", "author", "isbn"}:
            options["sort"] = str(report_options["sort"])
    if not any(bool(options[key]) for key in ("substances", "foreign_agents", "combined", "summary")):
        raise ValueError("Для Excel-отчёта должен быть выбран хотя бы один лист.")

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if progress_cb:
        progress_cb(85, "Создание выбранных списков точных совпадений")

    report_kwargs = {
        "deduplicate": bool(options["deduplicate"]),
        "sort_by": str(options["sort"]),
    }
    substance_records, substance_grouped = _report_records(results, SOURCE_SUBSTANCES, **report_kwargs)
    foreign_agent_records, foreign_agent_grouped = _report_records(results, SOURCE_FOREIGN_AGENTS, **report_kwargs)
    combined_records, combined_grouped = _report_records(results, **report_kwargs)

    workbook = Workbook()
    sheet_created = False
    if options["substances"]:
        _add_results_sheet(
            workbook,
            title="Вещества",
            headers=SUBSTANCE_MATCH_HEADERS,
            records=substance_records,
            grouped_results=substance_grouped,
            row_builder=_substance_match_row,
            table_name="SubstanceMatchesTable",
            table_style="TableStyleMedium4",
            row_fill_color="E2F0D9",
            widths={
                "A": 7,
                "B": 34,
                "C": 52,
                "D": 24,
                "E": 24,
                "F": 42,
                "G": 24,
                "H": 48,
                "I": 34,
                "J": 22,
                "K": 12,
            },
            active=not sheet_created,
            cancel_cb=cancel_cb,
        )
        sheet_created = True
    if options["foreign_agents"]:
        _add_results_sheet(
            workbook,
            title="Иностранные агенты",
            headers=FOREIGN_AGENT_MATCH_HEADERS,
            records=foreign_agent_records,
            grouped_results=foreign_agent_grouped,
            row_builder=_foreign_agent_match_row,
            table_name="ForeignAgentMatchesTable",
            table_style="TableStyleMedium9",
            row_fill_color="DDEBF7",
            widths={
                "A": 7,
                "B": 34,
                "C": 52,
                "D": 24,
                "E": 24,
                "F": 42,
                "G": 24,
                "H": 24,
                "I": 38,
                "J": 16,
                "K": 48,
                "L": 28,
                "M": 18,
                "N": 24,
                "O": 34,
                "P": 22,
                "Q": 12,
            },
            active=not sheet_created,
            cancel_cb=cancel_cb,
        )
        sheet_created = True
    if options["combined"]:
        _add_results_sheet(
            workbook,
            title="Все совпадения",
            headers=COMBINED_MATCH_HEADERS,
            records=combined_records,
            grouped_results=combined_grouped,
            row_builder=_combined_match_row,
            table_name="AllMatchesTable",
            table_style="TableStyleMedium2",
            row_fill_color="FFF2CC",
            widths={
                "A": 7,
                "B": 34,
                "C": 52,
                "D": 24,
                "E": 24,
                "F": 42,
                "G": 24,
                "H": 28,
                "I": 34,
                "J": 42,
            },
            active=not sheet_created,
            cancel_cb=cancel_cb,
        )
        sheet_created = True

    # Пограничные совпадения всегда выводятся отдельно: они видны оператору,
    # но никогда не попадают в списки подтверждённых совпадений и в метки.
    if (
        _add_review_sheet(
            workbook,
            results,
            active=not sheet_created,
            cancel_cb=cancel_cb,
        )
        is not None
    ):
        sheet_created = True
    if options["summary"]:
        _add_summary_sheet(workbook, summary, active=not sheet_created)
        sheet_created = True

    workbook.active = 0
    _matcher.atomic_write_via_path(output_path, workbook.save)
    summary.output_file = str(output_path)
    if progress_cb:
        progress_cb(
            92,
            f"Excel-отчёт создан: {output_path.name}. "
            f"Вещества: {len(substance_records):,}; "
            f"иноагенты: {len(foreign_agent_records):,}",
        )
    return output_path
