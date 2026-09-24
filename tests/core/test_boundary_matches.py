import tempfile
import unittest
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from irbis_control.core.matcher import (
    EXTRA_MATCH_RULES,
    SOURCE_SUBSTANCES,
    ComparisonSummary,
    DatabaseIndex,
    DatabaseRecord,
    ExcelEntry,
    ForeignAgentEntry,
    MatchResult,
    _deduplicate_cross_sheet_entries,
    _detect_header,
    _make_entry,
    _read_xlsx_entries,
    build_markers_by_record,
    compare_and_export,
    compare_database_records,
    compare_files,
    compare_foreign_agents,
    database_record_from_tag_values,
    export_results,
    normalize_author,
    normalize_publication_year,
    normalize_title,
    parse_match_rule,
    primary_author_contributor,
    read_foreign_agent_entries,
)


class BoundaryMatchTests(unittest.TestCase):
    def test_repeated_initial_before_full_name_is_the_same_author(self) -> None:
        variants = ("Фаулз Д. Джон", "Фаулз Дж. Джон", "Фаулз, Джон", "Фаулз Джон")

        self.assertEqual({"фаулз джон"}, {normalize_author(value) for value in variants})
        record = DatabaseRecord(record_number=1, authors=[variants[0]])
        matches = DatabaseIndex([record]).match_author_value(variants[1])

        self.assertEqual(1, len(matches))
        self.assertEqual(100.0, matches[0][2])

    def test_full_author_name_matches_in_natural_and_catalogue_order(self) -> None:
        for catalogue_name, natural_name in (
            ("Фаулз Джон", "Джон Фаулз"),
            ("Фаулз Джон Роберт", "Джон Роберт Фаулз"),
        ):
            with self.subTest(catalogue_name=catalogue_name):
                record = DatabaseRecord(record_number=1, authors=[catalogue_name])
                matches = DatabaseIndex([record]).match_author_value(natural_name)

                self.assertEqual(1, len(matches))
                self.assertEqual(100.0, matches[0][2])

    def test_isbn_10_and_isbn_13_are_equivalent(self) -> None:
        record = DatabaseRecord(record_number=1, isbns=["978-0-306-40615-7"])
        entry = ExcelEntry(1, "books.xlsx", "Книги", 2, isbn="0-306-40615-2")

        result = DatabaseIndex([record]).match(entry, True, True, False, 90)[0]

        self.assertEqual("Совпадение", result.status)
        self.assertEqual("ISBN", result.method)

    def test_isbn_with_conflicting_title_and_author_requires_review(self) -> None:
        record = DatabaseRecord(
            record_number=1,
            isbns=["978-0-306-40615-7"],
            titles=["Другая книга"],
            authors=["Петров Петр"],
        )
        entry = ExcelEntry(
            1,
            "books.xlsx",
            "Книги",
            2,
            isbn="978-0-306-40615-7",
            title="Коллекционер",
            author="Фаулз Джон",
        )

        result = DatabaseIndex([record]).match(entry, True, True, False, 90)[0]

        self.assertEqual("Возможное совпадение", result.status)
        self.assertEqual(95.0, result.confidence)
        self.assertIn("противоречием", result.method)

    def test_fuzzy_title_and_exact_author_allow_automatic_marker(self) -> None:
        record = DatabaseRecord(record_number=1, titles=["Коллекционер"], authors=["Фаулз Джон"])
        entry = ExcelEntry(
            1,
            "books.xlsx",
            "Книги",
            2,
            title="Колекционер",
            author="Фаулз, Джон",
        )

        result = DatabaseIndex([record]).match(entry, True, True, True, 92)[0]

        self.assertEqual("Совпадение", result.status)
        self.assertEqual(100.0, result.confidence)
        self.assertIn("точный автор", result.method)
        self.assertEqual({1: [(333, "^AIII")]}, build_markers_by_record([result]))

    def test_fuzzy_title_marks_all_copies_of_the_same_book(self) -> None:
        records = [
            DatabaseRecord(record_number=1, titles=["Коллекционер"], authors=["Фаулз Джон"]),
            DatabaseRecord(record_number=2, titles=["Коллекционер"], authors=["Фаулз Джон"]),
        ]
        entry = ExcelEntry(
            1,
            "books.xlsx",
            "Книги",
            2,
            title="Колекционер",
            author="Фаулз, Джон",
        )

        results = DatabaseIndex(records).match(entry, True, True, True, 92)

        self.assertEqual(2, len(results))
        self.assertTrue(all(result.status == "Совпадение" for result in results))
        self.assertEqual({1: [(333, "^AIII")], 2: [(333, "^AIII")]}, build_markers_by_record(results))

    def test_fuzzy_series_title_does_not_auto_confirm_wrong_volume(self) -> None:
        records = [
            DatabaseRecord(1, titles=["Эпоха мертвых. Москва"], authors=["Круз А. Ю. Андрей Юрьевич"]),
            DatabaseRecord(2, titles=["Эпоха мертвых. Начало"], authors=["Круз А. Ю. Андрей Юрьевич"]),
            DatabaseRecord(3, titles=["Эпоха мертвых. Прорыв"], authors=["Круз А. Ю. Андрей Юрьевич"]),
        ]
        entry = ExcelEntry(
            1,
            "books.xlsx",
            "Книги",
            2,
            title="Эпоха Мертвых : [18+] / Андрей Круз Начало",
            author="Круз, Андрей (1964-)",
        )

        results = DatabaseIndex(records).match(entry, True, True, True, 92)
        by_record = {result.database.record_number: result for result in results}

        self.assertEqual("Совпадение", by_record[2].status)
        self.assertEqual("Возможное совпадение", by_record[1].status)
        self.assertEqual("Возможное совпадение", by_record[3].status)

    def test_fuzzy_title_with_partial_author_requires_review(self) -> None:
        record = DatabaseRecord(record_number=1, titles=["Коллекционер"], authors=["Фаулз Д."])
        entry = ExcelEntry(
            1,
            "books.xlsx",
            "Книги",
            2,
            title="Колекционер",
            author="Фаулз Джон",
        )

        result = DatabaseIndex([record]).match(entry, True, True, True, 92)[0]

        self.assertEqual("Возможное совпадение", result.status)
        self.assertEqual({}, build_markers_by_record([result]))

    def test_publisher_legal_form_does_not_prevent_exact_publication_match(self) -> None:
        record = DatabaseRecord(
            record_number=1,
            titles=["Коллекционер"],
            publication=["^CИздательство Эксмо^D2024"],
        )
        entry = ExcelEntry(
            1,
            "books.xlsx",
            "Книги",
            2,
            title="Коллекционер",
            publisher="ООО «Эксмо»",
            year="2024",
        )

        result = DatabaseIndex([record]).match(entry, True, True, False, 90)[0]

        self.assertEqual("Совпадение", result.status)
        self.assertEqual("Название + издательство + год", result.method)

    def test_translator_in_field_702_is_not_compared_as_author(self) -> None:
        record = database_record_from_tag_values(
            1,
            [
                (200, "^AКоллекционер"),
                (700, "^AФаулз^BДж.^GДжон"),
                (702, "^AБессмертной^BИ.М.^4пер."),
            ],
        )
        translator_entry = ExcelEntry(
            1,
            "books.xlsx",
            "Книги",
            2,
            title="Коллекционер",
            author="Бессмертной И.М.",
        )
        primary_entry = replace(translator_entry, author="Фаулз Джон")

        translator = DatabaseIndex([record]).match(translator_entry, True, True, False, 90)[0]
        primary = DatabaseIndex([record]).match(primary_entry, True, True, False, 90)[0]

        self.assertEqual(["Фаулз Дж. Джон"], record.authors)
        self.assertEqual("Не найдено", translator.status)
        self.assertEqual({}, build_markers_by_record([translator]))
        self.assertEqual("Совпадение", primary.status)
        self.assertEqual({1: [(333, "^AIII")]}, build_markers_by_record([primary]))

    def test_non_author_roles_in_700_and_701_are_not_compared_as_authors(self) -> None:
        record = database_record_from_tag_values(
            1,
            [
                (200, "^AКнига"),
                (700, "^AРедакторов^BР.Р.^4340 редактор"),
                (701, "^AХудожников^BХ.Х.^4040 художник"),
                (701, "^AИллюстраторов^BИ.И.^4440 иллюстратор"),
                (701, "^AАвторов^BА.А.^4070 автор"),
                (701, "^AСоавторов^BС.С.^4aut"),
            ],
        )

        self.assertEqual(["Авторов А.А.", "Соавторов С.С."], record.authors)
        for excluded in ("Редакторов Р.Р.", "Художников Х.Х.", "Иллюстраторов И.И."):
            entry = ExcelEntry(1, "books.xlsx", "Книги", 2, title="Книга", author=excluded)
            self.assertEqual("Не найдено", DatabaseIndex([record]).match(entry, True, True, False, 90)[0].status)

    def test_tag_values_preserve_full_irbis_record_for_review(self) -> None:
        record = database_record_from_tag_values(
            17,
            [(10, "^A978-5-17-000000-0"), (200, "^AКнига"), (700, "^AАвтор^BА.А.")],
        )

        self.assertEqual(
            "#010: ^A978-5-17-000000-0\n#200: ^AКнига\n#700: ^AАвтор^BА.А.",
            record.raw_record,
        )

    def test_embedded_titles_and_authors_from_field_922_are_searchable(self) -> None:
        record = database_record_from_tag_values(
            1,
            [
                (200, "^AОкно с видом на площадь"),
                (700, "^AУитни^BФ.^GФиллис"),
                (922, "^CКимоно^FПэрис Д.^GДжон Пэрис"),
                (922, "^CКимоно^Uда/1"),
            ],
        )
        entry = ExcelEntry(
            1,
            "books.xlsx",
            "Книги",
            2,
            author="Пэрис Джон",
            title="Кимоно",
        )

        result = DatabaseIndex([record]).match(entry, False, True, False, 90)[0]

        self.assertEqual(["Окно с видом на площадь", "Кимоно"], record.titles)
        self.assertEqual(["Уитни Ф. Филлис", "Джон Пэрис"], record.authors)
        self.assertEqual("Совпадение", result.status)
        self.assertEqual("Название и автор", result.method)

    def test_repeated_titles_and_authors_are_all_preserved(self) -> None:
        record = database_record_from_tag_values(
            1,
            [
                (200, "^AПервое название^AВторое название"),
                (200, "^AТретье название"),
                (700, "^AПервый^BА."),
                (701, "^AВторой^BБ."),
            ],
        )

        self.assertEqual(["Первое название", "Второе название", "Третье название"], record.titles)
        self.assertEqual(["Первый А.", "Второй Б."], record.authors)

    def test_first_registry_contributor_is_used_as_primary_author(self) -> None:
        self.assertEqual(
            "Брэдбери, Рэй",
            primary_author_contributor("Брэдбери, Рэй Бабенко, Виталий Тимофеевич"),
        )
        self.assertEqual(
            "Мураками Харуки",
            primary_author_contributor("Мураками Харуки Чинарева, Юлия"),
        )
        record = DatabaseRecord(1, titles=["451 по Фаренгейту"], authors=["Брэдбери Р. Рэй"])
        entry = ExcelEntry(
            1,
            "books.xlsx",
            "Книги",
            2,
            author="Брэдбери, Рэй Бабенко, Виталий Тимофеевич",
            title="451 по Фаренгейту",
        )

        result = DatabaseIndex([record]).match(entry, False, True, False, 90)[0]

        self.assertEqual("Совпадение", result.status)
        self.assertEqual("Название и автор", result.method)

    def test_new_rules_match_independently_and_can_be_disabled(self) -> None:
        record = database_record_from_tag_values(
            1,
            [
                (10, "^A9780306406157"),
                (200, "^AТестовая книга"),
                (700, "^AИванов^BИ.И."),
                (210, "^CАСТ^D2024"),
            ],
        )
        entry = ExcelEntry(
            1,
            "books.xlsx",
            "Книги",
            2,
            author="Иванов И.И.",
            title="Тестовая книга",
            isbn="9780306406157",
            publisher="«АСТ»",
            year="2024 г.",
        )
        index = DatabaseIndex([record])
        for key, (label, required) in EXTRA_MATCH_RULES.items():
            with self.subTest(rule=key):
                result = index.match(entry, False, False, False, 90, {key: True})[0]
                self.assertEqual(label, result.method)
                weak = key in {
                    "title_year",
                    "title_publisher",
                    "author_year",
                    "author_publisher",
                    "author_publisher_year",
                }
                self.assertEqual("Возможное совпадение" if weak else "Совпадение", result.status)
                self.assertEqual(not weak, 1 in build_markers_by_record([result]))
                disabled = index.match(entry, False, False, False, 90, {key: False})
                self.assertEqual({}, build_markers_by_record(disabled))
                for missing in required:
                    incomplete = replace(entry, **{missing: ""})
                    result = index.match(incomplete, False, False, False, 90, {key: True})
                    self.assertEqual({}, build_markers_by_record(result))

    def test_rule_parser_normalizes_order_aliases_and_rejects_errors(self) -> None:
        for label, expected in (
            ("Год издания + Издатель + Заглавие", "title_publisher_year"),
            ("publisher, ISBN", "isbn_publisher"),
            ("Название + Автор", "use_title_fallback"),
            ("исбн", "use_isbn_matching"),
            ("год＋название", "title_year"),
            ("Автор + издательство + год", "author_publisher_year"),
        ):
            self.assertEqual(expected, parse_match_rule(label))
        for invalid in (
            "",
            "Название +",
            "Издательство + год",
            "Автор",
            "Название",
            "Название + Название",
            "ISBN + ИСБН",
            "Название + цена",
        ):
            with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                parse_match_rule(invalid)
        for key, (label, _fields) in EXTRA_MATCH_RULES.items():
            self.assertEqual(key, parse_match_rule(label))

    def test_isbn_combination_works_without_title_and_requires_metadata(self) -> None:
        record = DatabaseRecord(1, isbns=["9780306406157"], publication=["^CАСТ^D2024"])
        entry = ExcelEntry(1, "books.xlsx", "Книги", 2, isbn="9780306406157", year="2024")
        index = DatabaseIndex([record])
        result = index.match(entry, False, False, False, 90, {"isbn_year": True})[0]
        self.assertEqual("ISBN + год", result.method)
        self.assertEqual("Совпадение", result.status)
        for changed in (replace(entry, year="2023"), replace(entry, isbn="9780306406158")):
            self.assertEqual("Не найдено", index.match(changed, False, False, False, 90, {"isbn_year": True})[0].status)

    def test_placeholder_publisher_never_confirms_a_match(self) -> None:
        for publisher in ("нет", "[б. и.]", "не указано", "n/a"):
            entry = ExcelEntry(1, "books.xlsx", "Книги", 2, title="Книга", publisher=publisher, year="2024")
            record = DatabaseRecord(1, titles=[entry.title], publication=[f"^C{publisher}^D2024"])
            results = DatabaseIndex([record]).match(entry, False, False, False, 90, {"title_publisher_year": True})
            self.assertEqual({}, build_markers_by_record(results))

    def test_author_publication_rule_without_title_or_isbn_requires_review(self) -> None:
        records = [
            DatabaseRecord(1, authors=["Иванов И.И."], publication=["^CАСТ^D2024"]),
            DatabaseRecord(2, authors=["Петров П.П."], publication=["^CАСТ^D2024"]),
            DatabaseRecord(3, authors=["Иванов И.И."], publication=["^CЭксмо^D2024"]),
            DatabaseRecord(4, authors=["Иванов И.И."], publication=["^CАСТ^D2023"]),
        ]
        entry = ExcelEntry(1, "books.xlsx", "Книги", 2, author="Иванов И.И.", publisher="АСТ", year="2024")
        rules = {"author_publisher_year": True}
        results = DatabaseIndex(records).match(entry, False, False, False, 90, rules)
        self.assertEqual([1], [result.database.record_number for result in results])
        self.assertEqual("Возможное совпадение", results[0].status)
        self.assertEqual({}, build_markers_by_record(results))
        with (
            patch("irbis_control.core.matcher.read_excel_entries", return_value=([entry], [])),
            patch("irbis_control.core.matcher.read_foreign_agent_entries", return_value=([], [])),
        ):
            _, summary = compare_database_records(
                records, ["books.xlsx"], use_isbn_matching=False, use_title_fallback=False, match_rules=rules
            )
        self.assertEqual(1, summary.review_rows)
        self.assertEqual(0, summary.matched_excel_rows)

    def test_confirmed_rule_wins_over_author_review(self) -> None:
        record = DatabaseRecord(1, titles=["Книга"], authors=["Иванов И.И."], publication=["^CАСТ^D2024"])
        entry = ExcelEntry(
            1, "books.xlsx", "Книги", 2, title="Книга", author="Иванов И.И.", publisher="АСТ", year="2024"
        )
        results = DatabaseIndex([record]).match(
            entry, False, False, False, 90, {"author_publisher_year": True, "title_publisher_year": True}
        )
        self.assertEqual("Название + издательство + год", results[0].method)
        self.assertIn(1, build_markers_by_record(results))

    def test_strong_rule_wins_over_review_and_no_duplicate_record(self) -> None:
        record = DatabaseRecord(
            1, titles=["Тестовая книга", "Тестовая книга"], authors=["Иванов И.И."], publication=["^CАСТ^D2024"]
        )
        entry = ExcelEntry(
            1, "books.xlsx", "Книги", 2, title=record.main_title, author="Иванов И.И.", publisher="АСТ", year="2024"
        )
        rules = {"title_year": True, "title_author_publisher_year": True}
        results = DatabaseIndex([record]).match(entry, False, False, False, 90, rules)
        self.assertEqual(1, len(results))
        self.assertEqual("Название + автор + издательство + год", results[0].method)
        self.assertEqual(100.0, results[0].confidence)
        # Включённое прежнее правило по автору не должно ослабляться правилом ручной проверки.
        result = DatabaseIndex([record]).match(entry, False, True, False, 90, {"title_year": True})[0]
        self.assertEqual("Название и автор", result.method)
        self.assertEqual("Совпадение", result.status)

    def test_publisher_and_year_distinguish_editions_without_author(self) -> None:
        records = [
            DatabaseRecord(1, titles=["Тестовая книга"], publication=["^CАСТ^D2024"]),
            DatabaseRecord(2, titles=["Тестовая книга"], publication=["^CАСТ^D2023"]),
            DatabaseRecord(3, titles=["Тестовая книга"], publication=["^CЭксмо^D2024"]),
            DatabaseRecord(4, titles=["Тестовая книга"], publication=["^CАСТ", "^D2024"]),
        ]
        entry = ExcelEntry(1, "books.xlsx", "Книги", 2, title="Тестовая книга", publisher="АСТ", year="2024")
        results = DatabaseIndex(records).match(entry, False, False, False, 90, {"title_publisher_year": True})
        self.assertEqual([1], [result.database.record_number for result in results])

    def test_new_rules_reject_missing_database_fields_and_conflicting_author(self) -> None:
        entry = ExcelEntry(
            1, "books.xlsx", "Книги", 2, title="Тестовая книга", author="Иванов И.И.", publisher="АСТ", year="2024"
        )
        for publication, authors in [("^CАСТ", []), ("^D2024", []), ("^CАСТ^D2024", ["Петров П.П."])]:
            with self.subTest(publication=publication, authors=authors):
                record = DatabaseRecord(1, titles=[entry.title], authors=authors, publication=[publication])
                results = DatabaseIndex([record]).match(entry, False, False, False, 90, {"title_publisher_year": True})
                self.assertEqual({}, build_markers_by_record(results))

    def test_new_rules_do_not_treat_ambiguous_year_as_exact(self) -> None:
        record = DatabaseRecord(1, titles=["Тестовая книга"], publication=["^CАСТ^D2024"])
        for year in ("2023–2024", "2024?", "", "нет"):
            entry = ExcelEntry(1, "books.xlsx", "Книги", 2, title=record.main_title, publisher="АСТ", year=year)
            results = DatabaseIndex([record]).match(entry, False, False, False, 90, {"title_publisher_year": True})
            self.assertEqual({}, build_markers_by_record(results))

    def test_excel_publisher_and_year_headers_and_cross_sheet_editions(self) -> None:
        rows = [("Название", "Издательство", "Год издания"), ("Книга", "АСТ", 2024.0)]
        _, mapping, headers = _detect_header(rows)
        entry = _make_entry(1, Path("books.xlsx"), "Книги", 2, rows[1], mapping, headers)
        self.assertEqual(("АСТ", "2024"), (entry.publisher, entry.year))
        other_year = replace(entry, entry_id=2, sheet_name="Лист2", year="2023")
        other_publisher = replace(entry, entry_id=3, sheet_name="Лист2", publisher="Эксмо")
        duplicate = replace(entry, entry_id=4, sheet_name="Лист2")
        entries, skipped = _deduplicate_cross_sheet_entries([entry, other_year, other_publisher, duplicate])
        self.assertEqual([1, 2, 3], [item.entry_id for item in entries])
        self.assertEqual(1, skipped)
        with self.assertRaises(ValueError):
            _detect_header([("Издательство", "Год")])

    def test_rules_reach_direct_and_txt_comparison_and_summary(self) -> None:
        record = DatabaseRecord(1, titles=["Книга"], publication=["^CАСТ^D2024"])
        entry = ExcelEntry(1, "books.xlsx", "Книги", 2, title="Книга", publisher="АСТ", year="2024")
        options = {
            "use_isbn_matching": False,
            "use_title_fallback": False,
            "match_rules": {"title_publisher_year": True},
        }
        for with_foreign in (False, True):
            foreign = [self._foreign_entry(1, "Петров Петр Петрович")] if with_foreign else []
            with (
                patch("irbis_control.core.matcher.read_excel_entries", return_value=([entry], [])),
                patch("irbis_control.core.matcher.read_foreign_agent_entries", return_value=(foreign, [])),
                patch("irbis_control.core.matcher.parse_database", return_value=[record]),
            ):
                for results, summary in (
                    compare_database_records([record], ["books.xlsx"], **options),
                    compare_files(["database.txt"], ["books.xlsx"], **options),
                ):
                    self.assertEqual(1, summary.matched_excel_rows)
                    self.assertEqual(1, summary.exact_title_rows)
                    self.assertIn(1, build_markers_by_record(results))

    def test_xlsx_reader_loads_publication_fields(self) -> None:
        workbook = Workbook()
        workbook.active.append(["Title", "Publisher", "Year"])
        workbook.active.append(["Книга", "АСТ", 2024])
        with patch("irbis_control.core.matcher._load_workbook_quiet", return_value=workbook):
            entries, warnings = _read_xlsx_entries(Path("books.xlsx"), 1)
        self.assertFalse(warnings)
        self.assertEqual(("АСТ", "2024"), (entries[0].publisher, entries[0].year))

    def test_real_source_publisher_header_and_year_prefixes_are_supported(self) -> None:
        rows = [
            ("Автор", "Заглавие", "Изд-во", "Год выпуска"),
            (
                "Кучерская, Майя Александровна",
                "Тётя Мотя : роман : [16+]",
                "АСТ, Редакция Елены Шубиной",
                "cop. 2023",
            ),
        ]
        _, mapping, headers = _detect_header(rows)
        entry = _make_entry(1, Path("books.xlsx"), "Книги", 2, rows[1], mapping, headers)

        self.assertEqual("АСТ, Редакция Елены Шубиной", entry.publisher)
        self.assertEqual("2023", normalize_publication_year(entry.year))
        self.assertEqual("2026", normalize_publication_year("печ. 2026"))
        self.assertEqual("2023", normalize_publication_year("сор. 2023"))
        for ambiguous in ("2023–2024", "2024?", "2026-"):
            self.assertEqual("", normalize_publication_year(ambiguous))

        record = DatabaseRecord(
            1,
            titles=["Тетя Мотя"],
            authors=["Кучерская М.А."],
            publication=["^CРедакция Елены Шубиной^D2023"],
        )
        result = DatabaseIndex([record]).match(
            entry,
            False,
            False,
            False,
            90,
            {"title_author_publisher_year": True},
        )[0]
        self.assertEqual("Совпадение", result.status)
        self.assertEqual("Название + автор + издательство + год", result.method)

    def test_title_that_is_also_a_bibliographic_word_remains_searchable(self) -> None:
        self.assertEqual("текст", normalize_title("Текст : роман : [18+]"))
        self.assertEqual("рассказы", normalize_title("Рассказы"))
        self.assertNotEqual(normalize_title("Я (не) робот"), normalize_title("Я робот"))
        self.assertNotEqual(normalize_title("Избранные рассказы"), normalize_title("Избранные романы"))

    def test_database_uses_961_primary_author_when_700_701_are_missing(self) -> None:
        record = database_record_from_tag_values(
            1,
            [
                (200, "^AСад смерти"),
                (961, "^ZДА^AСандему^BМ.^GМаргит"),
                (961, "^4340 ред.^AРедактор^BР."),
            ],
        )

        self.assertEqual(["Сандему М. Маргит"], record.authors)
        entry = ExcelEntry(1, "books.xlsx", "Книги", 2, author="Сандему, Маргит", title="Сад смерти")
        result = DatabaseIndex([record]).match(entry, False, True, False, 90)[0]
        self.assertEqual("Совпадение", result.status)

    def test_exact_title_and_matching_first_author_initial_are_automatic(self) -> None:
        record = DatabaseRecord(1, titles=["Общее название"], authors=["Иванов И."])
        entry = ExcelEntry(
            1,
            "books.xlsx",
            "Книги",
            2,
            author="Иванов Иван Петрович",
            title="Общее название",
        )

        result = DatabaseIndex([record]).match(entry, False, True, False, 90)[0]

        self.assertEqual("Совпадение", result.status)
        self.assertEqual("Название и автор (сокращённое имя)", result.method)
        self.assertEqual({1: [(333, "^AIII")]}, build_markers_by_record([result]))

    def test_substance_report_contains_match_reason_and_source_location(self) -> None:
        result = MatchResult(
            status="Совпадение",
            method="Название и автор",
            confidence=100.0,
            excel=ExcelEntry(1, "source.xlsx", "Книги", 2703, title="Тётя Мотя"),
            database=DatabaseRecord(record_number=40, titles=["Тетя Мотя"]),
            source_type=SOURCE_SUBSTANCES,
            matched_value="Тётя Мотя",
        )
        summary = ComparisonSummary(
            database_file="Тест",
            excel_files=["source.xlsx"],
            database_records=1,
            database_records_with_isbn=0,
            excel_rows=1,
            matched_excel_rows=1,
            unmatched_excel_rows=0,
            result_rows=1,
            exact_isbn_rows=0,
            exact_title_rows=1,
            probable_rows=0,
        )

        output = BytesIO()
        with patch(
            "irbis_control.core.matcher.atomic_write_via_path", side_effect=lambda _path, writer: writer(output)
        ):
            export_results("report.xlsx", [result], summary, report_options={"substances": True})
        output.seek(0)
        workbook = load_workbook(output, read_only=True)
        try:
            sheet = workbook["Вещества"]
            self.assertEqual("Почему добавлено", sheet["H1"].value)
            self.assertEqual("Название и автор — Тётя Мотя", sheet["H2"].value)
            self.assertEqual("source.xlsx", sheet["I2"].value)
            self.assertEqual("Книги", sheet["J2"].value)
            self.assertEqual("2703", sheet["K2"].value)
        finally:
            workbook.close()

    def test_new_rule_reaches_report_only_export_without_writing_files(self) -> None:
        record = DatabaseRecord(1, titles=["Тестовая книга"], publication=["^CАСТ^D2024"])
        entry = ExcelEntry(1, "books.xlsx", "Книги", 2, title=record.main_title, publisher="АСТ", year="2024")
        output = BytesIO()
        with (
            patch("irbis_control.core.matcher.read_excel_entries", return_value=([entry], [])),
            patch("irbis_control.core.matcher.read_foreign_agent_entries", return_value=([], [])),
            patch("irbis_control.core.matcher.parse_database", return_value=[record]),
            patch("irbis_control.core.matcher.atomic_write_via_path", side_effect=lambda _path, writer: writer(output)),
            patch("irbis_control.core.matcher.Path.mkdir"),
        ):
            results, summary = compare_and_export(
                "database.txt",
                ["books.xlsx"],
                "report.xlsx",
                "modified.txt",
                use_isbn_matching=False,
                use_title_fallback=False,
                match_rules={"title_publisher_year": True},
                report_options={"enabled": True, "substances": True, "report_only": True},
            )
        output.seek(0)
        workbook = load_workbook(output, read_only=True)
        try:
            cells = [value for sheet in workbook for row in sheet.iter_rows(values_only=True) for value in row]
            self.assertTrue(
                any(
                    isinstance(value, str) and value.startswith(EXTRA_MATCH_RULES["title_publisher_year"][0])
                    for value in cells
                )
            )
            self.assertEqual(1, summary.matched_excel_rows)
            self.assertEqual(0, summary.modified_database_records)
        finally:
            workbook.close()

    @staticmethod
    def _foreign_entry(entry_id: int, name: str) -> ForeignAgentEntry:
        return ForeignAgentEntry(
            entry_id=entry_id,
            source_file="foreign.xlsx",
            sheet_name="Реестр",
            row_number=entry_id + 1,
            registry_number=str(entry_id),
            name=name,
            agent_type="Физическое лицо",
        )

    def test_single_initial_with_multiple_people_requires_review(self) -> None:
        record = DatabaseRecord(record_number=10, authors=["Иванов И."])
        entries = [
            self._foreign_entry(1, "Иванов Иван Петрович"),
            self._foreign_entry(2, "Иванов Илья Сергеевич"),
        ]

        results = compare_foreign_agents([record], entries)

        self.assertEqual(2, len(results))
        self.assertTrue(all(result.status == "Возможное совпадение" for result in results))
        self.assertTrue(all(result.confidence == 90.0 for result in results))
        self.assertEqual({}, build_markers_by_record(results))

    def test_single_initial_is_automatic_for_one_registry_person(self) -> None:
        record = DatabaseRecord(record_number=11, authors=["Иванов И."])

        results = compare_foreign_agents([record], [self._foreign_entry(1, "Иванов Иван Петрович")])

        self.assertEqual("Совпадение", results[0].status)
        self.assertEqual(100.0, results[0].confidence)
        self.assertEqual({11: [(333, "^AI^@ИВАНОВ ИВАН ПЕТРОВИЧ")]}, build_markers_by_record(results))

    def test_foreign_agents_use_the_same_author_comparison_as_substances(self) -> None:
        record = DatabaseRecord(record_number=12, titles=["Тестовая книга"], authors=["Иванов И."])
        substance = ExcelEntry(
            entry_id=1,
            source_file="substances.xlsx",
            sheet_name="Книги",
            row_number=2,
            author="Иванов Иван Петрович",
            title="Тестовая книга",
        )

        substance_result = DatabaseIndex([record]).match(substance, True, True, False, 90)[0]
        foreign_result = compare_foreign_agents([record], [self._foreign_entry(1, "Иванов Иван Петрович")])[0]

        self.assertEqual(substance_result.status, foreign_result.status)
        self.assertEqual(substance_result.confidence, foreign_result.confidence)

    def test_foreign_agent_publication_list_preserves_each_book_row(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Список Книг"
        sheet.append(["Перечень изданий иностранных агентов"])
        sheet.append(["Дата обновления списка:"])
        sheet.append(["№ в реестре", "Автор", "Роль относительно произведения", "ISBN", "Заглавие"])
        sheet.append(["123", "Берсенева, Анна", "Автор", "978-5-00-000001-1", "Первая книга"])
        sheet.append(["123", "Берсенева, Анна", "Автор", "978-5-00-000002-8", "Вторая книга"])

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as temp_dir:
            source = Path(temp_dir) / "publication-foreign-agent.xlsx"
            workbook.save(source)
            entries, warnings = read_foreign_agent_entries(source)

        self.assertEqual([], warnings)
        self.assertEqual(2, len(entries))
        self.assertEqual("Берсенева, Анна", entries[0].name)
        self.assertEqual("Физическое лицо", entries[0].agent_type)
        self.assertEqual("Первая книга", entries[0].title)
        self.assertEqual("978-5-00-000001-1", entries[0].isbn)

        record = DatabaseRecord(record_number=1, authors=["Берсенева Анна"])
        results = compare_foreign_agents([record], entries)
        result = next(item for item in results if item.method == "Реестр иностранных агентов: Автор")
        self.assertEqual("Совпадение", result.status)
        self.assertEqual(
            {1: [(333, "^AI^@БЕРСЕНЕВА АННА")]},
            build_markers_by_record([result]),
        )

    def test_foreign_agent_publication_is_matched_by_isbn_even_for_non_author_role(self) -> None:
        entry = ForeignAgentEntry(
            entry_id=1,
            source_file="publication-foreign-agent.xlsx",
            sheet_name="Список Книг",
            row_number=10,
            registry_number="777",
            name="Иванова Мария Петровна",
            agent_type="Физическое лицо",
            title="Другая книга",
            isbn="978-5-04-244068-7",
            role="пер.",
        )
        record = DatabaseRecord(
            record_number=77,
            isbns=["978-5-04-244068-7"],
            titles=["Другая книга"],
            authors=["Петров Петр"],
        )

        results = compare_foreign_agents([record], [entry])
        publication = next(item for item in results if item.method == "Список изданий иноагентов: ISBN")

        self.assertEqual("Совпадение", publication.status)
        self.assertEqual(
            {77: [(333, "^AI^@ИВАНОВА МАРИЯ ПЕТРОВНА")]},
            build_markers_by_record([publication]),
        )

    def test_foreign_agent_publication_organization_uses_organization_marker(self) -> None:
        entry = ForeignAgentEntry(
            entry_id=1,
            source_file="publication-foreign-agent.xlsx",
            sheet_name="Список Книг",
            row_number=4,
            registry_number="860",
            name="Общество с ограниченной ответственностью «Собеседник-Медиа»",
            agent_type="Организация",
            title="Филворды",
            isbn="978-5-04-244068-7",
            role="изд.",
        )
        record = DatabaseRecord(record_number=5, isbns=["978-5-04-244068-7"], titles=["Филворды"])

        publication = next(
            item
            for item in compare_foreign_agents([record], [entry])
            if item.method == "Список изданий иноагентов: ISBN"
        )

        self.assertEqual(
            {5: [(333, "^AO^@ОБЩЕСТВО С ОГРАНИЧЕННОЙ ОТВЕТСТВЕННОСТЬЮ «СОБЕСЕДНИК-МЕДИА»")]},
            build_markers_by_record([publication]),
        )

    def test_exact_title_without_author_requires_review(self) -> None:
        record = DatabaseRecord(
            record_number=20,
            titles=["Принцип сперматозоида"],
            authors=["Петров П.П."],
        )
        entry = ExcelEntry(
            entry_id=1,
            source_file="substances.xlsx",
            sheet_name="Книги",
            row_number=2,
            title="Принцип сперматозоида",
        )

        results = DatabaseIndex([record]).match(entry, True, True, False, 90)

        self.assertEqual("Возможное совпадение", results[0].status)
        self.assertEqual("Только название", results[0].method)
        self.assertEqual({}, build_markers_by_record(results))

    def test_exact_title_and_author_remain_confirmed(self) -> None:
        record = DatabaseRecord(
            record_number=21,
            titles=["Принцип сперматозоида"],
            authors=["Петров П.П."],
        )
        entry = ExcelEntry(
            entry_id=1,
            source_file="substances.xlsx",
            sheet_name="Книги",
            row_number=2,
            author="Петров П.П.",
            title="Принцип сперматозоида",
        )

        results = DatabaseIndex([record]).match(entry, True, True, False, 90)

        self.assertEqual("Совпадение", results[0].status)
        self.assertEqual(100.0, results[0].confidence)
        self.assertIn(21, build_markers_by_record(results))

    def test_confidence_below_100_is_never_eligible_for_marker(self) -> None:
        result = MatchResult(
            status="Совпадение",
            method="Название и автор",
            confidence=99.0,
            excel=ExcelEntry(1, "source.xlsx", "Лист1", 2),
            database=DatabaseRecord(record_number=30),
            source_type=SOURCE_SUBSTANCES,
        )

        self.assertEqual({}, build_markers_by_record([result]))

    def test_excel_report_contains_review_sheet_with_reason_and_source(self) -> None:
        result = MatchResult(
            status="Возможное совпадение",
            method="Только название",
            confidence=75.0,
            excel=ExcelEntry(1, "source.xlsx", "Книги", 12, title="Книга"),
            database=DatabaseRecord(record_number=40, titles=["Книга"]),
            note="Точно совпало только название",
            source_type=SOURCE_SUBSTANCES,
            matched_value="Книга",
        )
        summary = ComparisonSummary(
            database_file="Тест",
            excel_files=["source.xlsx"],
            database_records=1,
            database_records_with_isbn=0,
            excel_rows=1,
            matched_excel_rows=0,
            unmatched_excel_rows=1,
            result_rows=1,
            exact_isbn_rows=0,
            exact_title_rows=0,
            probable_rows=1,
            review_rows=1,
        )

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as temp_dir:
            output = Path(temp_dir) / "report.xlsx"
            export_results(
                output,
                [result],
                summary,
                report_options={"substances": True},
            )
            workbook = load_workbook(output, read_only=True)
            try:
                self.assertIn("Требует проверки", workbook.sheetnames)
                sheet = workbook["Требует проверки"]
                self.assertEqual("Точно совпало только название", sheet["D2"].value)
                self.assertEqual("source.xlsx", sheet["I2"].value)
            finally:
                workbook.close()

    def test_excel_report_groups_review_candidates_by_record_and_source(self) -> None:
        record = DatabaseRecord(record_number=40, source_record_number=40, titles=["Книга"])
        results = [
            MatchResult(
                status="Возможное совпадение",
                method="Приблизительно по названию и автору",
                confidence=confidence,
                excel=ExcelEntry(index, "source.xlsx", "Книги", index + 10, title=title),
                database=record,
                note=f"Вариант {index}",
                source_type=SOURCE_SUBSTANCES,
                matched_value=title,
            )
            for index, (title, confidence) in enumerate((("Книга 1", 95.0), ("Книга 2", 97.0)), start=1)
        ]
        summary = ComparisonSummary(
            database_file="Тест",
            excel_files=["source.xlsx"],
            database_records=1,
            database_records_with_isbn=0,
            excel_rows=2,
            matched_excel_rows=0,
            unmatched_excel_rows=2,
            result_rows=2,
            exact_isbn_rows=0,
            exact_title_rows=0,
            probable_rows=2,
            review_rows=2,
        )

        with tempfile.TemporaryDirectory(dir=Path.cwd()) as temp_dir:
            output = Path(temp_dir) / "report.xlsx"
            export_results(output, results, summary, report_options={"substances": True, "summary": True})
            workbook = load_workbook(output, read_only=True)
            try:
                review_sheet = workbook["Требует проверки"]
                self.assertEqual(2, review_sheet.max_row)
                self.assertIn("Книга 1", review_sheet["E2"].value)
                self.assertIn("Книга 2", review_sheet["E2"].value)
                summary_values = dict(workbook["Сводка"].iter_rows(min_row=2, values_only=True))
                self.assertEqual(1, summary_values["Записей для ручной проверки"])
                self.assertEqual(2, summary_values["Вариантов совпадений для ручной проверки"])
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
