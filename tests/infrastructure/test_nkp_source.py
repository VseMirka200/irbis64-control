from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from irbis_control.infrastructure import nkp_source


def _xlsx_bytes(path: Path) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Список Книг"
    ws.append(["Автор", "Заглавие", "ISBN"])
    ws.append(["Иванов И. И.", "Тестовая книга", "978-5-00-000000-1"])
    wb.save(path)
    return path.read_bytes()


def test_fetch_registry_follows_xlsx_link(tmp_path, monkeypatch) -> None:
    workbook_path = tmp_path / "source.xlsx"
    workbook = _xlsx_bytes(workbook_path)
    page = b'<html><a href="/files/publication-drugs.xlsx">download</a></html>'

    def fake_request(url: str, *, timeout: int = 25):
        if url.endswith("publication-drugs.xlsx"):
            return workbook, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", url
        return page, "text/html; charset=utf-8", "https://nkp.rsl.ru/drug-literature-recommendations"

    monkeypatch.setattr(nkp_source, "_request", fake_request)
    monkeypatch.setattr(nkp_source, "_render_dom", lambda _url: None)

    result = nkp_source.fetch_nkp_drug_registry(tmp_path / "cache")
    assert not result.used_cache
    assert result.path.is_file()
    wb = load_workbook(result.path, read_only=True, data_only=True)
    try:
        assert wb.active["B2"].value == "Тестовая книга"
    finally:
        wb.close()


def test_fetch_registry_prefers_full_official_download(tmp_path, monkeypatch) -> None:
    workbook_path = tmp_path / "source.xlsx"
    workbook = _xlsx_bytes(workbook_path)
    requested: list[str] = []

    def fake_request(url: str, *, timeout: int = 25):
        requested.append(url)
        if url == nkp_source.NKP_DRUG_DOWNLOAD_URL:
            return workbook, "application/octet-stream", url
        raise AssertionError(f"Unexpected fallback request: {url}")

    monkeypatch.setattr(nkp_source, "_request", fake_request)

    result = nkp_source.fetch_nkp_drug_registry(tmp_path / "cache")

    assert not result.used_cache
    assert result.source_url == nkp_source.NKP_DRUG_DOWNLOAD_URL
    assert requested == [nkp_source.NKP_DRUG_DOWNLOAD_URL]


def test_fetch_registry_can_build_xlsx_from_rendered_table(tmp_path, monkeypatch) -> None:
    page = b"<html><body>spa shell</body></html>"
    rendered = """
    <html><body><table>
      <tr><th>Автор</th><th>Заглавие</th><th>ISBN</th></tr>
      <tr><td>Петров П. П.</td><td>Книга с сайта</td><td>978-5-00-000000-2</td></tr>
      <tr><td>Сидоров С. С.</td><td>Ещё книга</td><td>978-5-00-000000-3</td></tr>
    </table></body></html>
    """

    monkeypatch.setattr(
        nkp_source,
        "_request",
        lambda _url, timeout=25: (page, "text/html", "https://nkp.rsl.ru/drug-literature-recommendations"),
    )
    monkeypatch.setattr(nkp_source, "_render_dom", lambda _url: rendered)

    result = nkp_source.fetch_nkp_drug_registry(tmp_path / "cache")
    assert not result.used_cache
    wb = load_workbook(result.path, read_only=True, data_only=True)
    try:
        rows = list(wb.active.iter_rows(values_only=True))
        assert rows[1][1] == "Книга с сайта"
        assert rows[2][1] == "Ещё книга"
    finally:
        wb.close()


def test_fetch_registry_uses_valid_cache_when_site_is_down(tmp_path, monkeypatch) -> None:
    cache = tmp_path / "cache"
    cache.mkdir()
    cached = cache / nkp_source.CACHE_FILENAME
    _xlsx_bytes(cached)

    def fail_request(_url: str, *, timeout: int = 25):
        raise OSError("offline")

    monkeypatch.setattr(nkp_source, "_request", fail_request)
    monkeypatch.setattr(nkp_source, "_render_dom", lambda _url: None)

    result = nkp_source.fetch_nkp_drug_registry(cache)
    assert result.used_cache
    assert result.path == cached
    assert "последняя сохранённая копия" in result.detail


def test_fetch_registry_reports_direct_download_failure(tmp_path, monkeypatch) -> None:
    def fake_request(url: str, *, timeout: int = 25):
        if url == nkp_source.NKP_DRUG_DOWNLOAD_URL:
            raise OSError("certificate verify failed")
        return b"<html><body>spa shell</body></html>", "text/html", nkp_source.NKP_DRUG_PAGE_URL

    monkeypatch.setattr(nkp_source, "_request", fake_request)
    monkeypatch.setattr(nkp_source, "_render_dom", lambda _url: None)
    monkeypatch.setattr(nkp_source, "resource_path", lambda *_parts: tmp_path / "missing.xlsx")

    try:
        nkp_source.fetch_nkp_drug_registry(tmp_path / "cache")
    except RuntimeError as exc:
        message = str(exc)
    else:
        raise AssertionError("Ожидалась ошибка загрузки реестра")

    assert "запрос выгрузки" in message
    assert nkp_source.NKP_DRUG_DOWNLOAD_URL in message
    assert "OSError: certificate verify failed" in message


def test_fetch_foreign_agents_registry_follows_xlsx_link(tmp_path, monkeypatch) -> None:
    source = tmp_path / "foreign-source.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Список Книг"
    ws.append(["Список изданий, выпущенных иностранными агентами"])
    ws.append(["Дата обновления списка:"])
    ws.append(
        ["№ в реестре", "Автор", "Роль относительно произведения", "Рег. №", "ISBN", "Заглавие", "Выходные данные"]
    )
    ws.append(["743", "Тестовый Автор", "Автор", "1", "978-5-00-000000-1", "Тестовая книга", "Москва, 2026"])
    wb.save(source)
    workbook = source.read_bytes()
    page = b'<html><a href="/files/publication-foreign-agent.xlsx">download</a></html>'

    def fake_request(url: str, *, timeout: int = 25):
        if url.endswith("publication-foreign-agent.xlsx"):
            return workbook, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", url
        return page, "text/html; charset=utf-8", nkp_source.NKP_FOREIGN_AGENTS_PAGE_URL

    monkeypatch.setattr(nkp_source, "_request", fake_request)
    monkeypatch.setattr(nkp_source, "_render_dom", lambda _url: None)

    result = nkp_source.fetch_nkp_foreign_agents_registry(tmp_path / "cache")
    assert not result.used_cache
    assert result.path.name == nkp_source.FOREIGN_AGENTS_CACHE_FILENAME
    loaded = load_workbook(result.path, read_only=True, data_only=True)
    try:
        assert loaded.active["B4"].value == "Тестовый Автор"
        assert loaded.active["F4"].value == "Тестовая книга"
    finally:
        loaded.close()


def test_fetch_foreign_agents_registry_uses_own_cache(tmp_path, monkeypatch) -> None:
    cache = tmp_path / "cache"
    cache.mkdir()
    cached = cache / nkp_source.FOREIGN_AGENTS_CACHE_FILENAME
    wb = Workbook()
    ws = wb.active
    ws.append(["Автор", "Заглавие", "ISBN"])
    ws.append(["Автор из кэша", "Книга из кэша", "978-5-00-000000-1"])
    wb.save(cached)

    def fail_request(_url: str, *, timeout: int = 25):
        raise OSError("offline")

    monkeypatch.setattr(nkp_source, "_request", fail_request)
    monkeypatch.setattr(nkp_source, "_render_dom", lambda _url: None)

    result = nkp_source.fetch_nkp_foreign_agents_registry(cache)
    assert result.used_cache
    assert result.path == cached
    assert "последняя сохранённая копия" in result.detail


def test_fetch_registry_rejects_single_rendered_page_and_uses_full_bundled_copy(tmp_path, monkeypatch) -> None:
    page = b"<html><body>spa shell</body></html>"
    rows = ["<tr><th>Автор</th><th>Заглавие</th><th>ISBN</th></tr>"]
    rows.extend(
        f"<tr><td>Автор {index}</td><td>Книга {index}</td><td>978-5-00-{index:07d}-1</td></tr>" for index in range(50)
    )
    rendered = "<html><body><table>" + "".join(rows) + "</table></body></html>"

    bundled = tmp_path / "publication-drugs.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Список Книг"
    ws.append(["Автор", "Заглавие", "ISBN"])
    for index in range(300):
        ws.append([f"Полный автор {index}", f"Полная книга {index}", f"978-5-01-{index:07d}-1"])
    wb.save(bundled)

    monkeypatch.setattr(
        nkp_source,
        "_request",
        lambda _url, timeout=25: (page, "text/html", nkp_source.NKP_DRUG_PAGE_URL),
    )
    monkeypatch.setattr(nkp_source, "_render_dom", lambda _url: rendered)
    monkeypatch.setattr(nkp_source, "resource_path", lambda *_parts: bundled)

    result = nkp_source.fetch_nkp_drug_registry(tmp_path / "cache")

    assert result.used_cache
    assert "неполный результат" in result.detail.casefold() or "одну страницу" in result.detail.casefold()
    loaded = load_workbook(result.path, read_only=True, data_only=True)
    try:
        assert loaded.active.max_row == 301
        assert loaded.active["B2"].value == "Полная книга 0"
    finally:
        loaded.close()


def test_fetch_registry_discards_old_fifty_row_cache(tmp_path, monkeypatch) -> None:
    cache = tmp_path / "cache"
    cache.mkdir()
    cached = cache / nkp_source.CACHE_FILENAME
    wb = Workbook()
    ws = wb.active
    ws.append(["Автор", "Заглавие", "ISBN"])
    for index in range(50):
        ws.append([f"Старый автор {index}", f"Старая книга {index}", f"978-5-02-{index:07d}-1"])
    wb.save(cached)

    bundled = tmp_path / "publication-drugs.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Автор", "Заглавие", "ISBN"])
    for index in range(250):
        ws.append([f"Полный автор {index}", f"Полная книга {index}", f"978-5-03-{index:07d}-1"])
    wb.save(bundled)

    def fail_request(_url: str, *, timeout: int = 25):
        raise OSError("offline")

    monkeypatch.setattr(nkp_source, "_request", fail_request)
    monkeypatch.setattr(nkp_source, "_render_dom", lambda _url: None)
    monkeypatch.setattr(nkp_source, "resource_path", lambda *_parts: bundled)

    result = nkp_source.fetch_nkp_drug_registry(cache)

    assert result.used_cache
    loaded = load_workbook(result.path, read_only=True, data_only=True)
    try:
        assert loaded.active.max_row == 251
        assert loaded.active["B2"].value == "Полная книга 0"
    finally:
        loaded.close()
