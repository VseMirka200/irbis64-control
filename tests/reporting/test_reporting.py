import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

from irbis_control.core.matcher import compare_and_export
from irbis_control.core.models import DatabaseRecord, ExcelEntry, ForeignAgentEntry, MatchResult
from irbis_control.core.report_export import (
    COMBINED_MATCH_HEADERS,
    _combined_match_row,
    _foreign_agent_match_row,
    _substance_match_row,
)
from irbis_control.reporting.result_diff import _iter_xlsx_rows, compare_result_files, compare_text_files


class ReportingTests(unittest.TestCase):
    def test_xlsx_is_closed_when_reading_fails(self) -> None:
        workbook = Mock()
        workbook.sheetnames = ["Вещества"]
        workbook.__getitem__ = Mock(side_effect=ValueError("broken sheet"))
        with patch("irbis_control.reporting.result_diff._load_workbook_quiet", return_value=workbook):
            with self.assertRaisesRegex(ValueError, "broken sheet"):
                _iter_xlsx_rows(Path("report.xlsx"))
        workbook.close.assert_called_once()


    def test_foreign_agent_report_removes_duplicate_values_inside_cells(self) -> None:
        record = DatabaseRecord(394, authors=["Акунин Б. Борис"], titles=["Шпионский роман"])
        entries = [
            ForeignAgentEntry(
                entry_id=1,
                source_file="publication-foreign-agent.xlsx",
                sheet_name="Список Книг",
                row_number=1966,
                registry_number="743",
                name='Чхартишвили Григорий Шалвович "Борис Акунин"',
                agent_type="Физическое лицо",
                role="авт.",
            ),
            ForeignAgentEntry(
                entry_id=2,
                source_file="publication-foreign-agent.xlsx",
                sheet_name="Список Книг",
                row_number=2733,
                registry_number="743",
                name='Чхартишвили Григорий Шалвович "Борис Акунин"',
                agent_type="Физическое лицо",
                role="авт.",
            ),
            ForeignAgentEntry(
                entry_id=3,
                source_file="publication-foreign-agent.xlsx",
                sheet_name="Список Книг",
                row_number=2892,
                registry_number="743",
                name='Чхартишвили Григорий Шалвович "Борис Акунин"',
                agent_type="Физическое лицо",
                role="авт.",
            ),
        ]
        results = [
            MatchResult(
                status="Совпадение",
                method="Список изданий иноагентов: Название и автор",
                confidence=100.0,
                excel=ExcelEntry(index + 1, entry.source_file, entry.sheet_name, entry.row_number),
                database=record,
                note="Роль: авт.",
                source_type="Иностранные агенты",
                matched_value=entry.name,
                foreign_agent=entry,
            )
            for index, entry in enumerate(entries)
        ]
        results[1].method = "Список изданий иноагентов: ISBN"

        row = _foreign_agent_match_row(1, record, results)

        self.assertEqual(row[7], "Название и автор\nISBN")
        self.assertEqual(row[8], entries[0].name)
        self.assertEqual(row[9], "743")
        self.assertEqual(row[10], entries[0].name)
        self.assertEqual(row[11], "Физическое лицо")
        self.assertEqual(row[13], "Роль: авт.")
        self.assertEqual(row[14], "publication-foreign-agent.xlsx")
        self.assertEqual(row[15], "Список Книг")
        self.assertEqual(row[16], "1966, 2733, 2892")

    def test_substance_report_removes_duplicate_values_inside_cells(self) -> None:
        record = DatabaseRecord(10, authors=["Автор"], titles=["Книга"])
        results = [
            MatchResult(
                status="Совпадение",
                method="Название и автор",
                confidence=100.0,
                excel=ExcelEntry(1, "publication-drugs.xlsx", "Список", 100),
                database=record,
                matched_value="Книга",
            ),
            MatchResult(
                status="Совпадение",
                method="Название и автор",
                confidence=100.0,
                excel=ExcelEntry(2, "publication-drugs.xlsx", "Список", 120),
                database=record,
                matched_value="Книга",
            ),
        ]

        row = _substance_match_row(1, record, results)

        self.assertEqual(row[7], "Название и автор — Книга")
        self.assertEqual(row[8], "publication-drugs.xlsx")
        self.assertEqual(row[9], "Список")
        self.assertEqual(row[10], "100, 120")

    def test_combined_report_row_matches_header_count_and_deduplicates(self) -> None:
        record = DatabaseRecord(10, authors=["Автор"], titles=["Книга"])
        result = MatchResult(
            status="Совпадение",
            method="ISBN",
            confidence=100.0,
            excel=ExcelEntry(1, "list.xlsx", "Лист1", 2),
            database=record,
            source_type="Вещества",
            matched_value="9780000000000",
        )

        row = _combined_match_row(1, record, [result, result])

        self.assertEqual(len(row), len(COMBINED_MATCH_HEADERS))
        self.assertEqual(row[7], "Вещества")
        self.assertEqual(row[8], "ISBN")
        self.assertEqual(row[9], "9780000000000")

    def test_local_comparison_exports_reports_and_modified_database(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory)
            database = folder / "database.txt"
            original = "#010: ^A9780306406157\n#200: ^AТестовая книга\n#700: ^AИванов^BИ.И.\n*****\n"
            database.write_text(original, encoding="utf-8")
            source = folder / "books.xlsx"
            workbook = Workbook()
            workbook.active.append(["ISBN", "Название", "Автор"])
            workbook.active.append(["9780306406157", "Тестовая книга", "Иванов И.И."])
            workbook.save(source)
            workbook.close()

            reports = []
            for number in range(2):
                report = folder / f"report{number}.xlsx"
                modified = folder / f"modified{number}.txt"
                results, summary = compare_and_export(database, [source], report, modified)
                self.assertEqual(summary.modified_database_records, 1)
                self.assertTrue(any(result.status == "Совпадение" for result in results))
                self.assertTrue(modified.is_file())
                book = load_workbook(report, read_only=True)
                try:
                    self.assertIn("Вещества", book.sheetnames)
                    self.assertGreater(book["Вещества"].max_row, 1)
                finally:
                    book.close()
                reports.append(report)

            differences, summary = compare_result_files(*reports, folder / "changes.xlsx")
            self.assertEqual(differences, [])
            self.assertEqual(summary.total_changes, 0)
            differences, summary = compare_text_files(database, modified, folder / "changes.txt")
            self.assertEqual(summary.changed, 1)
            self.assertEqual(len(differences), 1)
            self.assertEqual(database.read_text(encoding="utf-8"), original)


if __name__ == "__main__":
    unittest.main()
